"""Export filtered search results to CSV or Excel.

Data transformation is Streamlit-independent so it can be unit tested.
UI rendering lives in ``streamlit_results.render_export_control``.
"""

from __future__ import annotations

import csv
import hashlib
import html
import io
import logging
import re
from dataclasses import dataclass
from datetime import date, datetime, timezone
from typing import Any, Iterable, Literal, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from rental_search_agent.analysis_view import match_strength
from rental_search_agent.display_format import (
    safe_http_url,
    score_to_pct,
    split_listing_address,
)
from rental_search_agent.geocoding import NEAREST_TRANSIT_LOCATION
from rental_search_agent.preference_criteria import AMENITY_FEATURES, match_amenity_feature

logger = logging.getLogger(__name__)

ExportFormat = Literal["xlsx", "csv"]
ExportScope = Literal["filtered", "selected", "page"]

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_HTML_TAG_RE = re.compile(r"<[^>]+>")
_UNSAFE_FILENAME_RE = re.compile(r"[^\w.\-]+", re.UNICODE)
_MAX_EXCEL_COL_WIDTH = 48
_CURRENCY_HEADERS = frozenset(
    {
        "Listing Price",
        "Strata / Maintenance Fee",
        "Property Tax",
    }
)
_PERCENT_HEADERS = frozenset(
    {
        "Overall Match (%)",
        "Structural Score (%)",
        "Proximity Score (%)",
        "Amenity Score (%)",
        "Semantic Score (%)",
    }
)
_TEXT_FORCE_HEADERS = frozenset(
    {
        "MLS Number",
        "Postal Code",
        "Unit Number",
    }
)

# Explicit column order for the primary results sheet / CSV header.
EXPORT_COLUMNS: tuple[str, ...] = (
    # Identity
    "MLS Number",
    "Full Address",
    "Street Address",
    "Unit Number",
    "City",
    "Province",
    "Postal Code",
    "Listing URL",
    "Source",
    "Brokerage",
    "Listing Agent",
    # Property details
    "Property Type",
    "MLS Category",
    "Ownership Type",
    "Listing Type",
    "Listing Price",
    "Price Display",
    "Bedrooms",
    "Bedrooms Display",
    "Has Den",
    "Bathrooms",
    "Interior Size (sq ft)",
    "Lot Size",
    "Stories",
    "Year Built",
    "Parking Spaces",
    "Parking Type",
    "Strata / Maintenance Fee",
    "Property Tax",
    "Listing Age",
    "Days on Market",
    "Open House",
    "Price Change",
    "Rank",
    # Amenities (value + evidence)
    "Balcony",
    "Balcony Evidence",
    "Gym",
    "Gym Evidence",
    "In-suite Laundry",
    "In-suite Laundry Evidence",
    "Laundry in Building",
    "Laundry in Building Evidence",
    "Walk-in Closet",
    "Walk-in Closet Evidence",
    "Ensuite Bathroom",
    "Ensuite Bathroom Evidence",
    "Parking",
    "Parking Evidence",
    "Storage",
    "Storage Evidence",
    "Air Conditioning",
    "Air Conditioning Evidence",
    "Pet Policy",
    "Pet Policy Evidence",
    "Dishwasher",
    "Dishwasher Evidence",
    "Swimming Pool",
    "Swimming Pool Evidence",
    "EV Charger",
    "EV Charger Evidence",
    "Elevator",
    "Elevator Evidence",
    "Furnished",
    "Furnished Evidence",
    "Den",
    "Den Evidence",
    "Listed Amenities",
    "Nearby Amenities",
    # Location / proximity
    "Latitude",
    "Longitude",
    "Nearest Transit Walking Distance (km)",
    "Nearest Transit Walking Time (min)",
    "Proximity Details",
    # Match analysis
    "Overall Match (%)",
    "Overall Match Label",
    "Structural Score (%)",
    "Proximity Score (%)",
    "Amenity Score (%)",
    "Semantic Score (%)",
    "Criteria Evaluated Count",
    "Criteria Total Count",
    "Standout Reasons",
    "Open Questions",
    "Unmet Criteria",
    # Listing text / AI
    "Original Listing Description",
    "AI Listing Summary",
    "AI Key Matches",
    "AI Key Gaps",
)

_AMENITY_EXPORT_MAP: tuple[tuple[str, str, str], ...] = (
    ("balcony", "Balcony", "Balcony Evidence"),
    ("gym", "Gym", "Gym Evidence"),
    ("laundry", "In-suite Laundry", "In-suite Laundry Evidence"),
    ("laundry_building", "Laundry in Building", "Laundry in Building Evidence"),
    ("walk_in_closet", "Walk-in Closet", "Walk-in Closet Evidence"),
    ("ensuite_bathroom", "Ensuite Bathroom", "Ensuite Bathroom Evidence"),
    ("parking", "Parking", "Parking Evidence"),
    ("storage", "Storage", "Storage Evidence"),
    ("ac", "Air Conditioning", "Air Conditioning Evidence"),
    ("pets", "Pet Policy", "Pet Policy Evidence"),
    ("dishwasher", "Dishwasher", "Dishwasher Evidence"),
    ("pool", "Swimming Pool", "Swimming Pool Evidence"),
    ("ev", "EV Charger", "EV Charger Evidence"),
    ("elevator", "Elevator", "Elevator Evidence"),
    ("furnished", "Furnished", "Furnished Evidence"),
    ("den", "Den", "Den Evidence"),
)

_AMENITY_BY_ID = {feat.id: feat for feat in AMENITY_FEATURES}

SCOPE_LABELS: dict[ExportScope, str] = {
    "filtered": "Current filtered results",
    "selected": "Selected properties",
    "page": "Current page",
}


@dataclass(frozen=True)
class ExportScopeOption:
    scope: ExportScope
    label: str
    count: int


@dataclass(frozen=True)
class PreparedExport:
    filename: str
    mime_type: str
    data: bytes
    cache_key: str
    row_count: int
    scope: ExportScope
    format: ExportFormat

    @property
    def size_bytes(self) -> int:
        return len(self.data)

    @property
    def size_label(self) -> str:
        n = self.size_bytes
        if n < 1024:
            return f"{n} B"
        if n < 1024 * 1024:
            return f"{n / 1024:.1f} KB"
        return f"{n / (1024 * 1024):.1f} MB"


def available_export_scopes(
    *,
    filtered_count: int,
    selected_count: int = 0,
    page_count: int | None = None,
    pagination_enabled: bool = False,
    selection_enabled: bool = False,
) -> list[ExportScopeOption]:
    """Return scopes the UI should offer. Only include what the app supports."""
    options = [
        ExportScopeOption(
            scope="filtered",
            label=f"{SCOPE_LABELS['filtered']} ({filtered_count})",
            count=filtered_count,
        )
    ]
    if selection_enabled and selected_count > 0:
        options.append(
            ExportScopeOption(
                scope="selected",
                label=f"{SCOPE_LABELS['selected']} ({selected_count})",
                count=selected_count,
            )
        )
    if pagination_enabled and page_count is not None:
        options.append(
            ExportScopeOption(
                scope="page",
                label=f"{SCOPE_LABELS['page']} ({page_count})",
                count=page_count,
            )
        )
    return options


def resolve_export_listings(
    *,
    scope: ExportScope,
    filtered_listings: Sequence[Mapping[str, Any]],
    selected_listings: Sequence[Mapping[str, Any]] | None = None,
    page_listings: Sequence[Mapping[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    """Resolve the listing set for the chosen scope. Preserves input order."""
    if scope == "selected":
        source = selected_listings or []
    elif scope == "page":
        source = page_listings or []
    else:
        source = filtered_listings
    return [dict(item) for item in source if isinstance(item, Mapping)]


def sanitize_export_text(value: Any) -> str:
    """Neutralize spreadsheet formula injection and strip HTML to plain text."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    text = str(value)
    text = html.unescape(text)
    text = _HTML_TAG_RE.sub(" ", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text).strip()
    if not text:
        return ""
    if text.startswith(_FORMULA_PREFIXES):
        return f"'{text}"
    return text


def flatten_list_value(value: Any) -> str:
    """Short lists → semicolon-separated text; complex structures → compact JSON-like text."""
    if value is None:
        return ""
    if isinstance(value, str):
        return sanitize_export_text(value)
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)):
        return "" if isinstance(value, float) and value != value else str(value)
    if isinstance(value, Mapping):
        parts = []
        for key, item in value.items():
            flat = flatten_list_value(item)
            if flat:
                parts.append(f"{key}: {flat}")
        return sanitize_export_text("; ".join(parts))
    if isinstance(value, (list, tuple, set)):
        parts = [flatten_list_value(item) for item in value]
        parts = [p for p in parts if p]
        return sanitize_export_text("; ".join(parts))
    return sanitize_export_text(value)


def _blank_row() -> dict[str, str]:
    return {col: "" for col in EXPORT_COLUMNS}


def _bool_export(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)) and value in (0, 1):
        return "Yes" if value else "No"
    text = str(value).strip().lower()
    if text in ("yes", "true", "y", "1"):
        return "Yes"
    if text in ("no", "false", "n", "0"):
        return "No"
    if text in ("unknown", "n/a", "na", "none"):
        return ""
    return ""


def _score_pct_cell(value: Any) -> str:
    pct = score_to_pct(value)
    return "" if pct is None else str(pct)


def _number_cell(value: Any) -> str:
    if value is None or value == "":
        return ""
    try:
        n = float(value)
    except (TypeError, ValueError):
        return ""
    if n != n or n in (float("inf"), float("-inf")):
        return ""
    if n == int(n):
        return str(int(n))
    return f"{n:g}"


def _parse_city_province(address: str | None, postal_code: str | None) -> tuple[str, str, str, str]:
    """Return (full_address, street, city, province) without inventing missing parts."""
    full = (address or "").strip()
    headline, locality = split_listing_address(full or None, postal_code)
    street = headline if full else ""
    city = ""
    province = ""
    if locality:
        # Typical: "Vancouver, British Columbia V6B 1A1" or "Vancouver, BC"
        loc_parts = [p.strip() for p in locality.split(",") if p.strip()]
        if loc_parts:
            city = loc_parts[0]
        if len(loc_parts) >= 2:
            rest = loc_parts[1]
            # Drop trailing postal from province token when glued together.
            rest = re.sub(
                r"\b[A-Za-z]\d[A-Za-z]\s?\d[A-Za-z]\d\b",
                "",
                rest,
            ).strip()
            province = rest
    return full, street, city, province


def _days_on_market(listing: Mapping[str, Any]) -> str:
    hours = listing.get("listing_age_hours")
    try:
        h = float(hours)
    except (TypeError, ValueError):
        return ""
    if h != h or h < 0:
        return ""
    return str(int(h // 24))


def _transit_walk_metrics(proximity: Any) -> tuple[str, str]:
    """Nearest-transit walking distance_km and duration_min when present."""
    if not isinstance(proximity, Mapping):
        return "", ""
    for key, val in proximity.items():
        raw = str(key or "")
        location, _, mode = raw.partition("|")
        loc = location.strip().lower()
        mode_l = mode.strip().lower()
        is_transit = (
            loc in {NEAREST_TRANSIT_LOCATION.lower(), "nearest transit", "transit"}
            or loc.startswith("nearest transit")
        )
        if not is_transit:
            continue
        if mode_l and mode_l not in ("walk", "walking"):
            continue
        if not isinstance(val, Mapping):
            continue
        dist = _number_cell(val.get("distance_km"))
        mins = _number_cell(val.get("duration_min"))
        return dist, mins
    return "", ""


def _proximity_details(proximity: Any) -> str:
    if not isinstance(proximity, Mapping) or not proximity:
        return ""
    parts: list[str] = []
    for key, val in proximity.items():
        location, _, mode = str(key).partition("|")
        label = location.strip() or "location"
        if mode.strip():
            label = f"{label} ({mode.strip()})"
        if val is None:
            parts.append(f"{label}: unavailable")
            continue
        if not isinstance(val, Mapping):
            parts.append(f"{label}: unavailable")
            continue
        bits = []
        mins = _number_cell(val.get("duration_min"))
        dist = _number_cell(val.get("distance_km"))
        if mins:
            bits.append(f"{mins} min")
        if dist:
            bits.append(f"{dist} km")
        parts.append(f"{label}: {' / '.join(bits)}" if bits else f"{label}: unavailable")
    return sanitize_export_text("; ".join(parts))


def _amenity_cells(listing: Mapping[str, Any]) -> dict[str, str]:
    out: dict[str, str] = {}
    for feat_id, value_col, evidence_col in _AMENITY_EXPORT_MAP:
        feat = _AMENITY_BY_ID.get(feat_id)
        if feat is None:
            out[value_col] = ""
            out[evidence_col] = ""
            continue
        result = match_amenity_feature(listing, feat)
        if result.status == "met":
            out[value_col] = sanitize_export_text(result.observed or "Yes")
            out[evidence_col] = sanitize_export_text(result.source or "Unknown")
        elif result.status == "unmet":
            out[value_col] = "No"
            out[evidence_col] = sanitize_export_text(result.source or "Unknown")
        else:
            # unknown — do not invent; leave value blank, evidence Unknown when description missing
            out[value_col] = ""
            out[evidence_col] = "Unknown" if result.source is None else sanitize_export_text(result.source)
    return out


def _checklist_items(listing: Mapping[str, Any], analysis: Mapping[str, Any] | None) -> list[dict]:
    breakdown = None
    if analysis and isinstance(analysis.get("score_breakdown"), Mapping):
        breakdown = analysis.get("score_breakdown")
    if breakdown is None:
        breakdown = listing.get("score_breakdown")
    if not isinstance(breakdown, Mapping):
        return []
    checklist = breakdown.get("checklist")
    if not isinstance(checklist, list):
        return []
    return [item for item in checklist if isinstance(item, Mapping)]


def _component_scores(
    listing: Mapping[str, Any], analysis: Mapping[str, Any] | None
) -> dict[str, Any]:
    breakdown = None
    if analysis and isinstance(analysis.get("score_breakdown"), Mapping):
        breakdown = analysis["score_breakdown"]
    if breakdown is None and isinstance(listing.get("score_breakdown"), Mapping):
        breakdown = listing["score_breakdown"]
    if not isinstance(breakdown, Mapping):
        return {}
    components = breakdown.get("components")
    return components if isinstance(components, Mapping) else {}


def _criteria_counts(checklist: Sequence[Mapping[str, Any]]) -> tuple[str, str]:
    if not checklist:
        return "", ""
    total = len(checklist)
    evaluated = sum(1 for item in checklist if str(item.get("status") or "") != "unknown")
    return str(evaluated), str(total)


def _criteria_labels(checklist: Sequence[Mapping[str, Any]], status: str) -> str:
    labels = []
    for item in checklist:
        if str(item.get("status") or "") != status:
            continue
        name = str(item.get("name") or item.get("label") or item.get("id") or "").strip()
        if name:
            labels.append(name)
    return sanitize_export_text("; ".join(labels))


def listing_to_export_row(
    listing: Mapping[str, Any],
    *,
    analysis: Mapping[str, Any] | None = None,
) -> dict[str, str]:
    """Convert one canonical listing (+ optional Analyze result) into an export row."""
    row = _blank_row()
    analysis = analysis if isinstance(analysis, Mapping) and "error" not in analysis else None

    full, street, city, province = _parse_city_province(
        listing.get("address") if isinstance(listing.get("address"), str) else None,
        listing.get("postal_code") if isinstance(listing.get("postal_code"), str) else None,
    )
    url = safe_http_url(listing.get("url") if isinstance(listing.get("url"), str) else None) or ""

    row["MLS Number"] = sanitize_export_text(listing.get("id"))
    row["Full Address"] = sanitize_export_text(full)
    row["Street Address"] = sanitize_export_text(street)
    row["Unit Number"] = ""  # not in canonical Listing model
    row["City"] = sanitize_export_text(city)
    row["Province"] = sanitize_export_text(province)
    row["Postal Code"] = sanitize_export_text(listing.get("postal_code"))
    row["Listing URL"] = sanitize_export_text(url)
    row["Source"] = sanitize_export_text(listing.get("source"))
    row["Brokerage"] = sanitize_export_text(listing.get("brokerage_name"))
    row["Listing Agent"] = sanitize_export_text(listing.get("agent_name"))

    row["Property Type"] = sanitize_export_text(listing.get("house_category"))
    row["MLS Category"] = sanitize_export_text(listing.get("property_category"))
    row["Ownership Type"] = sanitize_export_text(listing.get("ownership_category"))
    row["Listing Type"] = sanitize_export_text(listing.get("listing_type"))
    row["Listing Price"] = _number_cell(listing.get("price"))
    row["Price Display"] = sanitize_export_text(listing.get("price_display"))
    row["Bedrooms"] = _number_cell(listing.get("bedrooms"))
    row["Bedrooms Display"] = sanitize_export_text(
        listing.get("bedrooms_display") or listing.get("bedrooms")
    )
    row["Has Den"] = _bool_export(listing.get("has_den"))
    row["Bathrooms"] = _number_cell(listing.get("bathrooms"))
    row["Interior Size (sq ft)"] = _number_cell(listing.get("sqft"))
    row["Lot Size"] = sanitize_export_text(listing.get("lot_size"))
    row["Stories"] = _number_cell(listing.get("stories"))
    row["Year Built"] = ""  # not scraped today
    row["Parking Spaces"] = _number_cell(listing.get("parking_spaces"))
    row["Parking Type"] = sanitize_export_text(listing.get("parking_type"))
    row["Strata / Maintenance Fee"] = ""  # not scraped today
    row["Property Tax"] = ""  # not scraped today
    row["Listing Age"] = sanitize_export_text(listing.get("listing_age_display"))
    row["Days on Market"] = _days_on_market(listing)
    row["Open House"] = sanitize_export_text(listing.get("open_house"))
    row["Price Change"] = sanitize_export_text(listing.get("price_change_display"))
    row["Rank"] = _number_cell(listing.get("rank"))

    row.update(_amenity_cells(listing))
    row["Listed Amenities"] = sanitize_export_text(listing.get("ammenities"))
    row["Nearby Amenities"] = sanitize_export_text(listing.get("nearby_ammenities"))

    row["Latitude"] = _number_cell(listing.get("latitude"))
    row["Longitude"] = _number_cell(listing.get("longitude"))
    transit_dist, transit_min = _transit_walk_metrics(listing.get("proximity"))
    row["Nearest Transit Walking Distance (km)"] = transit_dist
    row["Nearest Transit Walking Time (min)"] = transit_min
    row["Proximity Details"] = _proximity_details(listing.get("proximity"))

    match_score = listing.get("match_score")
    if match_score is None:
        match_score = listing.get("semantic_score")
    if analysis and analysis.get("match_score_pct") is not None:
        overall_pct = score_to_pct(analysis.get("match_score_pct"))
    else:
        overall_pct = score_to_pct(match_score)
    row["Overall Match (%)"] = "" if overall_pct is None else str(overall_pct)
    label, _blurb = match_strength(overall_pct)
    row["Overall Match Label"] = sanitize_export_text(label) if overall_pct is not None else ""

    components = _component_scores(listing, analysis)
    row["Structural Score (%)"] = _score_pct_cell(components.get("structural"))
    row["Proximity Score (%)"] = _score_pct_cell(components.get("proximity"))
    row["Amenity Score (%)"] = _score_pct_cell(components.get("amenity"))
    row["Semantic Score (%)"] = _score_pct_cell(
        components.get("semantic") if components.get("semantic") is not None else listing.get("semantic_score")
    )

    checklist = _checklist_items(listing, analysis)
    evaluated, total = _criteria_counts(checklist)
    row["Criteria Evaluated Count"] = evaluated
    row["Criteria Total Count"] = total

    standout = ""
    open_questions = _criteria_labels(checklist, "unknown")
    unmet = _criteria_labels(checklist, "unmet")
    if analysis:
        standout = flatten_list_value(analysis.get("key_matches"))
        if analysis.get("key_gaps"):
            unmet = flatten_list_value(analysis.get("key_gaps")) or unmet
        row["AI Listing Summary"] = sanitize_export_text(analysis.get("ai_listing_summary"))
        row["AI Key Matches"] = flatten_list_value(analysis.get("key_matches"))
        row["AI Key Gaps"] = flatten_list_value(analysis.get("key_gaps"))
    else:
        # Prefer met checklist labels as standout when Analyze has not been run.
        standout = _criteria_labels(checklist, "met")
        row["AI Listing Summary"] = ""
        row["AI Key Matches"] = ""
        row["AI Key Gaps"] = ""
    row["Standout Reasons"] = standout
    row["Open Questions"] = open_questions
    row["Unmet Criteria"] = unmet
    row["Original Listing Description"] = sanitize_export_text(listing.get("description"))
    return row


def successful_analysis_by_id(
    analysis_by_id: Mapping[str, Mapping[str, Any]] | None,
) -> dict[str, Mapping[str, Any]]:
    """Keep only usable Analyze payloads (drop error stubs) for export + cache keys."""
    out: dict[str, Mapping[str, Any]] = {}
    for lid, payload in (analysis_by_id or {}).items():
        if not isinstance(payload, Mapping):
            continue
        if "error" in payload:
            continue
        key = str(lid).strip()
        if key:
            out[key] = payload
    return out


def listings_to_export_rows(
    listings: Sequence[Mapping[str, Any]],
    *,
    analysis_by_id: Mapping[str, Mapping[str, Any]] | None = None,
) -> list[dict[str, str]]:
    analysis_by_id = successful_analysis_by_id(analysis_by_id)
    rows: list[dict[str, str]] = []
    for listing in listings:
        lid = str(listing.get("id") or "")
        analysis = analysis_by_id.get(lid)
        rows.append(listing_to_export_row(listing, analysis=analysis))
    return rows


def format_active_filters(prefs: Mapping[str, Any] | None) -> str:
    """Readable summary of Search Preferences for Export Info."""
    prefs = prefs or {}
    parts: list[str] = []
    mapping = (
        ("location", "Location"),
        ("listing_type", "Listing type"),
        ("budget_max", "Budget max"),
        ("min_bedrooms", "Min bedrooms"),
        ("max_bedrooms", "Max bedrooms"),
        ("min_bathrooms", "Min bathrooms"),
        ("min_sqft", "Min size (sq ft)"),
        ("house_categories", "Property types"),
        ("proximity_preferences", "Proximity"),
        ("preferred_directions", "Facing"),
        ("qualitative_preferences", "Qualitative"),
    )
    for key, label in mapping:
        raw = prefs.get(key)
        if raw is None:
            continue
        text = str(raw).strip()
        if not text:
            continue
        if key == "preferred_directions":
            from rental_search_agent.preferred_directions import display_labels, parse_preferred_directions

            labels = display_labels(parse_preferred_directions(raw))
            if not labels:
                continue
            text = ", ".join(labels)
        parts.append(f"{label}: {sanitize_export_text(text)}")
    return "; ".join(parts)


def build_export_filename(
    *,
    scope: ExportScope,
    count: int,
    format: ExportFormat,
    when: date | None = None,
) -> str:
    day = (when or datetime.now(timezone.utc).date()).isoformat()
    scope_token = {
        "filtered": "filtered",
        "selected": "selected",
        "page": "page",
    }.get(scope, "filtered")
    raw = f"property_search_{scope_token}_{count}_{day}.{format}"
    return _UNSAFE_FILENAME_RE.sub("_", raw)


def export_cache_key(
    *,
    listings: Sequence[Mapping[str, Any]],
    scope: ExportScope,
    format: ExportFormat,
    sort_by: str | None,
    filters_text: str,
    analysis_ids: Iterable[str] | None = None,
) -> str:
    """Stable key so prepared bytes are reused across Streamlit reruns."""
    identity = []
    for item in listings:
        identity.append(
            (
                str(item.get("id") or ""),
                item.get("rank"),
                item.get("match_score"),
                item.get("semantic_score"),
                item.get("price"),
                bool(item.get("description")),
                bool(item.get("score_breakdown")),
            )
        )
    payload = {
        "identity": identity,
        "scope": scope,
        "format": format,
        "sort_by": sort_by or "",
        "filters": filters_text or "",
        "analysis_ids": sorted({str(x) for x in (analysis_ids or []) if x}),
        "columns": list(EXPORT_COLUMNS),
    }
    digest = hashlib.sha256(repr(payload).encode("utf-8")).hexdigest()
    return digest


def build_csv_bytes(rows: Sequence[Mapping[str, str]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.DictWriter(
        buffer,
        fieldnames=list(EXPORT_COLUMNS),
        extrasaction="ignore",
        lineterminator="\n",
        quoting=csv.QUOTE_MINIMAL,
    )
    writer.writeheader()
    for row in rows:
        writer.writerow({col: row.get(col, "") for col in EXPORT_COLUMNS})
    # utf-8-sig BOM for Excel compatibility
    return buffer.getvalue().encode("utf-8-sig")


def _style_header(ws: Worksheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _autosize_columns(ws: Worksheet, rows: Sequence[Mapping[str, str]]) -> None:
    for idx, header in enumerate(EXPORT_COLUMNS, start=1):
        max_len = len(header)
        for row in rows[:200]:
            val = row.get(header, "") or ""
            # Cap sample length so huge descriptions don't dominate width calc.
            max_len = max(max_len, min(len(val), 60))
        width = min(_MAX_EXCEL_COL_WIDTH, max(10, max_len + 2))
        ws.column_dimensions[get_column_letter(idx)].width = width


def _write_results_sheet(ws: Worksheet, rows: Sequence[Mapping[str, str]]) -> None:
    ws.append(list(EXPORT_COLUMNS))
    wrap_headers = {
        "Original Listing Description",
        "AI Listing Summary",
        "Proximity Details",
        "Standout Reasons",
        "Open Questions",
        "Unmet Criteria",
        "Listed Amenities",
        "Nearby Amenities",
        "AI Key Matches",
        "AI Key Gaps",
    }
    for row in rows:
        values = []
        for col in EXPORT_COLUMNS:
            raw = row.get(col, "")
            if col in _CURRENCY_HEADERS and raw != "":
                try:
                    values.append(float(raw))
                    continue
                except ValueError:
                    pass
            if col in _PERCENT_HEADERS and raw != "":
                try:
                    values.append(int(raw))
                    continue
                except ValueError:
                    pass
            if col in _TEXT_FORCE_HEADERS:
                values.append(str(raw) if raw != "" else "")
                continue
            values.append(raw)
        ws.append(values)
    _style_header(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for r_idx in range(2, len(rows) + 2):
        for c_idx, header in enumerate(EXPORT_COLUMNS, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            if header in wrap_headers:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if header in _CURRENCY_HEADERS and isinstance(cell.value, (int, float)):
                cell.number_format = '"$"#,##0'
            if header in _PERCENT_HEADERS and isinstance(cell.value, (int, float)):
                cell.number_format = "0"
            if header in _TEXT_FORCE_HEADERS and cell.value is not None:
                cell.number_format = "@"
    _autosize_columns(ws, rows)


def _write_export_info_sheet(
    ws: Worksheet,
    *,
    row_count: int,
    scope: ExportScope,
    sort_label: str | None,
    filters_text: str,
    exported_at: datetime | None = None,
) -> None:
    when = exported_at or datetime.now(timezone.utc)
    evidence_note = (
        "Evidence labels: MLS = structured listing field; Calculated = derived "
        "(e.g. travel time); Inferred = matched from description/amenity text; "
        "Unknown = not enough evidence to decide."
    )
    rows = [
        ("Export timestamp (UTC)", when.strftime("%Y-%m-%d %H:%M:%S UTC")),
        ("Number of exported properties", str(row_count)),
        ("Export scope", SCOPE_LABELS.get(scope, scope)),
        ("Active filters", filters_text or "None recorded"),
        ("Active sort order", (sort_label or "").strip() or "Default / unspecified"),
        ("Evidence label guide", evidence_note),
    ]
    ws.append(["Field", "Value"])
    for field, value in rows:
        ws.append([field, sanitize_export_text(value)])
    _style_header(ws)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 80
    for cell in ws["B"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def build_xlsx_bytes(
    rows: Sequence[Mapping[str, str]],
    *,
    scope: ExportScope,
    sort_label: str | None = None,
    filters_text: str = "",
    exported_at: datetime | None = None,
) -> bytes:
    wb = Workbook()
    ws_results = wb.active
    ws_results.title = "Search Results"
    _write_results_sheet(ws_results, rows)
    ws_info = wb.create_sheet("Export Info")
    _write_export_info_sheet(
        ws_info,
        row_count=len(rows),
        scope=scope,
        sort_label=sort_label,
        filters_text=filters_text,
        exported_at=exported_at,
    )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def prepare_export(
    listings: Sequence[Mapping[str, Any]],
    *,
    scope: ExportScope,
    format: ExportFormat,
    sort_by: str | None = None,
    sort_label: str | None = None,
    filters_text: str = "",
    analysis_by_id: Mapping[str, Mapping[str, Any]] | None = None,
    when: date | None = None,
) -> PreparedExport:
    """Build export bytes for the given listings and format."""
    usable_analysis = successful_analysis_by_id(analysis_by_id)
    rows = listings_to_export_rows(listings, analysis_by_id=usable_analysis)
    cache_key = export_cache_key(
        listings=listings,
        scope=scope,
        format=format,
        sort_by=sort_by,
        filters_text=filters_text,
        analysis_ids=usable_analysis.keys(),
    )
    filename = build_export_filename(
        scope=scope, count=len(rows), format=format, when=when
    )
    if format == "csv":
        data = build_csv_bytes(rows)
        mime = "text/csv"
    else:
        data = build_xlsx_bytes(
            rows,
            scope=scope,
            sort_label=sort_label or sort_by,
            filters_text=filters_text,
        )
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return PreparedExport(
        filename=filename,
        mime_type=mime,
        data=data,
        cache_key=cache_key,
        row_count=len(rows),
        scope=scope,
        format=format,
    )
