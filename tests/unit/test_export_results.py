"""Unit tests for search-results CSV/Excel export."""

from __future__ import annotations

import csv
import inspect
import io
from datetime import date
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook

from rental_search_agent.export_results import (
    EXPORT_COLUMNS,
    available_export_scopes,
    build_csv_bytes,
    build_export_filename,
    build_xlsx_bytes,
    export_cache_key,
    flatten_list_value,
    format_active_filters,
    listing_to_export_row,
    listings_to_export_rows,
    prepare_export,
    resolve_export_listings,
    sanitize_export_text,
    successful_analysis_by_id,
)
from rental_search_agent.streamlit_results import (
    render_export_control,
    render_search_results,
)


def _sample_listing(**overrides):
    base = {
        "id": "R1234567",
        "title": "Bright 2 bed",
        "url": "https://www.realtor.ca/listing/R1234567",
        "address": "123 Main St, Vancouver, British Columbia V6B1A1",
        "postal_code": "V6B 1A1",
        "price": 2800.0,
        "price_display": "$2,800/month",
        "bedrooms": 2,
        "has_den": True,
        "bedrooms_display": "2 + 1",
        "bathrooms": 2.0,
        "sqft": 950.0,
        "source": "Realtor.ca",
        "house_category": "Apartment",
        "property_category": "Single Family",
        "ownership_category": "Condominium",
        "listing_type": "for_rent",
        "latitude": 49.28,
        "longitude": -123.12,
        "parking_spaces": 1,
        "parking_type": "Underground",
        "lot_size": "0",
        "listing_age_hours": 48.0,
        "listing_age_display": "2 days ago",
        "brokerage_name": "Example Realty",
        "agent_name": "Ada Agent",
        "description": "Sunny balcony home with in-suite laundry and gym access.",
        "ammenities": "Balcony, Gym",
        "rank": 1,
        "match_score": 0.82,
        "semantic_score": 0.7,
        "score_breakdown": {
            "components": {
                "structural": 0.9,
                "proximity": 0.8,
                "amenity": 0.75,
                "semantic": 0.7,
            },
            "checklist": [
                {"id": "beds", "name": "Bedrooms", "status": "met", "source": "MLS"},
                {"id": "balcony", "name": "Balcony", "status": "met", "source": "Inferred"},
                {"id": "pets", "name": "Pet-friendly", "status": "unknown"},
                {"id": "budget", "name": "Budget", "status": "unmet", "source": "MLS"},
            ],
        },
        "proximity": {
            "nearest transit station|walk": {"duration_min": 6, "distance_km": 0.4},
            "Downtown|drive": {"duration_min": 18, "distance_km": 7.2},
        },
    }
    base.update(overrides)
    return base


class TestSanitizeAndFlatten:
    def test_formula_injection_neutralized(self):
        for raw in ("=1+1", "+cmd", "-2+3", "@SUM(A1)", "\t=HID", "\r=HID"):
            out = sanitize_export_text(raw)
            assert out.startswith("'")
            assert out[1:] == raw or out.endswith(raw.lstrip("\t\r"))

    def test_html_stripped_to_plain_text(self):
        assert "<" not in sanitize_export_text("<b>Nice</b> place")
        assert "Nice" in sanitize_export_text("<b>Nice</b> place")

    def test_lists_flattened(self):
        assert flatten_list_value(["a", "b", "c"]) == "a; b; c"
        assert flatten_list_value({"x": 1, "y": None}) == "x: 1"

    def test_nullish_values_blank(self):
        row = listing_to_export_row({"id": "x", "url": "https://www.realtor.ca/x", "address": "", "price": 0, "bedrooms": 0, "title": "t"})
        assert row["Year Built"] == ""
        assert row["Property Tax"] == ""
        assert row["AI Listing Summary"] == ""


class TestScopeResolution:
    def test_filtered_is_default_and_only_scope_without_selection_or_pagination(self):
        scopes = available_export_scopes(filtered_count=127)
        assert [s.scope for s in scopes] == ["filtered"]
        assert "127" in scopes[0].label

    def test_selected_and_page_only_when_supported(self):
        scopes = available_export_scopes(
            filtered_count=20,
            selected_count=8,
            page_count=10,
            pagination_enabled=True,
            selection_enabled=True,
        )
        assert [s.scope for s in scopes] == ["filtered", "selected", "page"]
        assert "8" in scopes[1].label
        assert "10" in scopes[2].label

    def test_resolve_selected_and_page(self):
        filtered = [_sample_listing(id="a"), _sample_listing(id="b"), _sample_listing(id="c")]
        selected = [filtered[0], filtered[2]]
        page = [filtered[1]]
        assert [x["id"] for x in resolve_export_listings(scope="filtered", filtered_listings=filtered)] == [
            "a",
            "b",
            "c",
        ]
        assert [x["id"] for x in resolve_export_listings(
            scope="selected", filtered_listings=filtered, selected_listings=selected
        )] == ["a", "c"]
        assert [x["id"] for x in resolve_export_listings(
            scope="page", filtered_listings=filtered, page_listings=page
        )] == ["b"]


class TestRowSchema:
    def test_explicit_column_order_and_expected_fields(self):
        row = listing_to_export_row(_sample_listing())
        assert list(row.keys()) == list(EXPORT_COLUMNS)
        assert row["MLS Number"] == "R1234567"
        assert row["Full Address"].startswith("123 Main St")
        assert row["City"] == "Vancouver"
        assert "British Columbia" in row["Province"]
        assert row["Postal Code"] == "V6B 1A1"
        assert row["Listing Price"] == "2800"
        assert row["Property Type"] == "Apartment"
        assert row["Overall Match (%)"] == "82"
        assert row["Overall Match Label"] == "Good match"
        assert row["Structural Score (%)"] == "90"
        assert row["Nearest Transit Walking Time (min)"] == "6"
        assert row["Nearest Transit Walking Distance (km)"] == "0.4"
        assert "balcony" in row["Original Listing Description"].lower()
        assert row["Criteria Evaluated Count"] == "3"
        assert row["Criteria Total Count"] == "4"
        assert "Pet-friendly" in row["Open Questions"]
        assert "Budget" in row["Unmet Criteria"]

    def test_analysis_fields_included_when_available(self):
        analysis = {
            "match_score_pct": 88,
            "key_matches": ["Near transit", "Has balcony"],
            "key_gaps": ["No pets"],
            "ai_listing_summary": "A bright apartment with balcony access.",
            "score_breakdown": {
                "components": {"structural": 0.9},
                "checklist": [],
            },
        }
        row = listing_to_export_row(_sample_listing(), analysis=analysis)
        assert row["Overall Match (%)"] == "88"
        assert row["AI Listing Summary"].startswith("A bright apartment")
        assert "Near transit" in row["Standout Reasons"]
        assert "No pets" in row["Unmet Criteria"]
        assert "Near transit" in row["AI Key Matches"]

    def test_amenity_evidence_columns(self):
        row = listing_to_export_row(_sample_listing())
        assert row["Balcony"] in ("Yes", "Balcony", "balcony") or row["Balcony"]
        assert row["Balcony Evidence"] in ("MLS", "Inferred", "Unknown")


class TestCsvAndExcel:
    def test_csv_utf8_bom_columns_and_unicode(self):
        listings = [
            _sample_listing(id="1", address="Rue Saint-Denis, Montréal, QC"),
            _sample_listing(id="2", description="Spacious unit — café nearby"),
        ]
        rows = listings_to_export_rows(listings)
        data = build_csv_bytes(rows)
        assert data.startswith(b"\xef\xbb\xbf")
        text = data.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        assert reader.fieldnames == list(EXPORT_COLUMNS)
        parsed = list(reader)
        assert len(parsed) == 2
        assert "Montréal" in parsed[0]["Full Address"]
        assert "café" in parsed[1]["Original Listing Description"]

    def test_excel_sheets_and_formatting(self):
        rows = listings_to_export_rows([_sample_listing(), _sample_listing(id="R999", rank=2)])
        data = build_xlsx_bytes(
            rows,
            scope="filtered",
            sort_label="Match",
            filters_text="Location: Vancouver",
        )
        wb = load_workbook(io.BytesIO(data))
        assert wb.sheetnames == ["Search Results", "Export Info"]
        ws = wb["Search Results"]
        assert ws["A1"].value == "MLS Number"
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref
        assert ws.max_row == 3  # header + 2
        info = wb["Export Info"]
        values = {info.cell(r, 1).value: info.cell(r, 2).value for r in range(2, info.max_row + 1)}
        assert values["Number of exported properties"] == "2"
        assert values["Export scope"] == "Current filtered results"
        assert "Vancouver" in values["Active filters"]
        assert "MLS" in values["Evidence label guide"]

    def test_formula_injection_in_csv_and_excel(self):
        listing = _sample_listing(description="=HYPERLINK(\"http://evil\")")
        rows = listings_to_export_rows([listing])
        csv_text = build_csv_bytes(rows).decode("utf-8-sig")
        assert "'=HYPERLINK" in csv_text or "''=HYPERLINK" in csv_text
        wb = load_workbook(io.BytesIO(build_xlsx_bytes(rows, scope="filtered")))
        desc_col = list(EXPORT_COLUMNS).index("Original Listing Description") + 1
        cell = wb["Search Results"].cell(row=2, column=desc_col).value
        assert str(cell).startswith("'")


class TestPrepareAndCache:
    def test_filename_sanitized(self):
        name = build_export_filename(
            scope="filtered", count=127, format="xlsx", when=date(2026, 9, 14)
        )
        assert name == "property_search_filtered_127_2026-09-14.xlsx"

    def test_prepare_export_csv_and_xlsx(self):
        listings = [_sample_listing(), _sample_listing(id="2", rank=2)]
        xlsx = prepare_export(listings, scope="filtered", format="xlsx", sort_label="Price")
        csv_prepared = prepare_export(listings, scope="filtered", format="csv")
        assert xlsx.filename.endswith(".xlsx")
        assert csv_prepared.filename.endswith(".csv")
        assert xlsx.row_count == 2
        assert csv_prepared.row_count == 2
        assert xlsx.size_bytes > 0

    def test_cache_key_changes_when_filters_or_data_change(self):
        listings = [_sample_listing()]
        k1 = export_cache_key(
            listings=listings,
            scope="filtered",
            format="xlsx",
            sort_by="match_score",
            filters_text="Location: Vancouver",
        )
        k2 = export_cache_key(
            listings=listings,
            scope="filtered",
            format="xlsx",
            sort_by="match_score",
            filters_text="Location: Burnaby",
        )
        k3 = export_cache_key(
            listings=[_sample_listing(id="other")],
            scope="filtered",
            format="xlsx",
            sort_by="match_score",
            filters_text="Location: Vancouver",
        )
        k4 = export_cache_key(
            listings=listings,
            scope="filtered",
            format="csv",
            sort_by="match_score",
            filters_text="Location: Vancouver",
        )
        assert k1 != k2
        assert k1 != k3
        assert k1 != k4
        assert k1 == export_cache_key(
            listings=listings,
            scope="filtered",
            format="xlsx",
            sort_by="match_score",
            filters_text="Location: Vancouver",
        )

    def test_cache_key_ignores_failed_analysis_entries(self):
        listings = [_sample_listing()]
        ok = {"match_score_pct": 80, "key_matches": ["balcony"], "key_gaps": []}
        mixed = {
            "R1234567": ok,
            "other": {"error": "boom"},
        }
        usable = successful_analysis_by_id(mixed)
        assert set(usable) == {"R1234567"}
        prepared = prepare_export(
            listings,
            scope="filtered",
            format="csv",
            analysis_by_id=mixed,
        )
        ui_key = export_cache_key(
            listings=listings,
            scope="filtered",
            format="csv",
            sort_by=None,
            filters_text="",
            analysis_ids=usable.keys(),
        )
        assert prepared.cache_key == ui_key
        # Including the failed id in the UI key must not happen — that was the bug.
        bad_ui_key = export_cache_key(
            listings=listings,
            scope="filtered",
            format="csv",
            sort_by=None,
            filters_text="",
            analysis_ids=mixed.keys(),
        )
        assert prepared.cache_key != bad_ui_key

    def test_format_active_filters(self):
        text = format_active_filters(
            {"location": "Vancouver", "min_bedrooms": "2", "budget_max": ""}
        )
        assert "Location: Vancouver" in text
        assert "Min bedrooms: 2" in text
        assert "Budget" not in text

    def test_format_active_filters_facing_labels(self):
        text = format_active_filters({"preferred_directions": "S,W"})
        assert "Facing: South, West" in text


class TestUiWiring:
    def test_toolbar_used_by_all_result_views(self):
        source = inspect.getsource(render_search_results)
        assert "render_results_toolbar" in source
        assert "table" in source and "map" in source and "grid" in source or "_render_results_grid" in source

    def test_export_control_is_in_toolbar_not_per_view(self):
        from rental_search_agent import streamlit_results as sr

        header_src = inspect.getsource(sr._render_results_header)
        assert "render_export_control" in header_src
        grid_src = inspect.getsource(sr._render_results_grid)
        table_src = inspect.getsource(sr._render_results_table)
        map_src = inspect.getsource(sr._render_map_panel)
        assert "render_export_control" not in grid_src
        assert "render_export_control" not in table_src
        assert "render_export_control" not in map_src

    def test_same_filtered_dataset_exported_regardless_of_view(self):
        listings = [_sample_listing(id="a"), _sample_listing(id="b")]
        # View mode is not an input to prepare_export — only the listing list matters.
        a = prepare_export(listings, scope="filtered", format="csv")
        b = prepare_export(listings, scope="filtered", format="csv")
        assert a.data == b.data
        assert a.cache_key == b.cache_key

    def test_disabled_export_when_no_results(self):
        fake_st = MagicMock()
        fake_st.session_state = {}
        with patch("rental_search_agent.streamlit_results.st", fake_st):
            render_export_control([], enabled=False)
        fake_st.button.assert_called()
        kwargs = fake_st.button.call_args.kwargs
        assert kwargs.get("disabled") is True
        fake_st.popover.assert_not_called()

    def test_prepare_error_keeps_results_and_shows_message(self):
        fake_st = MagicMock()
        fake_st.session_state = {
            "export_scope": "filtered",
            "export_format": "xlsx",
            "user_preferences": {},
            "analysis_result": {},
        }
        # Enter popover context
        fake_st.popover.return_value.__enter__ = MagicMock(return_value=MagicMock())
        fake_st.popover.return_value.__exit__ = MagicMock(return_value=False)
        fake_st.radio.side_effect = ["filtered", "xlsx"]
        fake_st.button.return_value = True  # Prepare export clicked
        fake_st.spinner.return_value.__enter__ = MagicMock(return_value=None)
        fake_st.spinner.return_value.__exit__ = MagicMock(return_value=False)

        with patch("rental_search_agent.streamlit_results.st", fake_st):
            with patch(
                "rental_search_agent.export_results.prepare_export",
                side_effect=RuntimeError("boom"),
            ):
                render_export_control([_sample_listing()], enabled=True)

        assert fake_st.session_state["export_error"] == (
            "We couldn’t prepare the export. Please try again."
        )
        assert "export_prepared" not in fake_st.session_state
        fake_st.error.assert_called()

    def test_rerun_reuses_prepared_bytes_without_regenerating(self):
        listings = [_sample_listing()]
        prepared = prepare_export(listings, scope="filtered", format="csv")
        fake_st = MagicMock()
        fake_st.session_state = {
            "export_scope": "filtered",
            "export_format": "csv",
            "user_preferences": {},
            "analysis_result": {},
            "export_prepared": {
                "cache_key": prepared.cache_key,
                "filename": prepared.filename,
                "mime_type": prepared.mime_type,
                "data": prepared.data,
                "row_count": prepared.row_count,
                "size_label": prepared.size_label,
            },
        }
        fake_st.popover.return_value.__enter__ = MagicMock(return_value=MagicMock())
        fake_st.popover.return_value.__exit__ = MagicMock(return_value=False)
        fake_st.radio.side_effect = ["filtered", "csv"]
        fake_st.button.return_value = False  # do not prepare again

        with patch("rental_search_agent.streamlit_results.st", fake_st):
            with patch(
                "rental_search_agent.export_results.prepare_export"
            ) as prepare_mock:
                render_export_control(listings, enabled=True)
                prepare_mock.assert_not_called()
        fake_st.download_button.assert_called()

    def test_filter_change_invalidates_prepared_export(self):
        listings = [_sample_listing()]
        prepared = prepare_export(
            listings,
            scope="filtered",
            format="csv",
            filters_text="Location: Vancouver",
        )
        fake_st = MagicMock()
        fake_st.session_state = {
            "export_scope": "filtered",
            "export_format": "csv",
            "user_preferences": {"location": "Burnaby"},
            "analysis_result": {},
            "export_prepared": {
                "cache_key": prepared.cache_key,
                "filename": prepared.filename,
                "mime_type": prepared.mime_type,
                "data": prepared.data,
                "row_count": prepared.row_count,
                "size_label": prepared.size_label,
            },
        }
        fake_st.popover.return_value.__enter__ = MagicMock(return_value=MagicMock())
        fake_st.popover.return_value.__exit__ = MagicMock(return_value=False)
        fake_st.radio.side_effect = ["filtered", "csv"]
        fake_st.button.return_value = False

        with patch("rental_search_agent.streamlit_results.st", fake_st):
            render_export_control(listings, enabled=True)
        assert "export_prepared" not in fake_st.session_state
        fake_st.download_button.assert_not_called()

    def test_failed_analysis_does_not_block_download_button(self):
        listings = [_sample_listing()]
        prepared = prepare_export(
            listings,
            scope="filtered",
            format="csv",
            analysis_by_id={"R1234567": {"error": "failed"}},
        )
        fake_st = MagicMock()
        fake_st.session_state = {
            "export_scope": "filtered",
            "export_format": "csv",
            "user_preferences": {},
            "analysis_result": {"R1234567": {"error": "failed"}},
            "export_prepared": {
                "cache_key": prepared.cache_key,
                "filename": prepared.filename,
                "mime_type": prepared.mime_type,
                "data": prepared.data,
                "row_count": prepared.row_count,
                "size_label": prepared.size_label,
            },
        }
        fake_st.popover.return_value.__enter__ = MagicMock(return_value=MagicMock())
        fake_st.popover.return_value.__exit__ = MagicMock(return_value=False)
        fake_st.radio.side_effect = ["filtered", "csv"]
        fake_st.button.return_value = False

        with patch("rental_search_agent.streamlit_results.st", fake_st):
            render_export_control(listings, enabled=True)
        fake_st.download_button.assert_called()

    def test_zero_results_uses_shared_toolbar(self):
        from rental_search_agent.streamlit_landing import render_zero_results

        source = inspect.getsource(render_zero_results)
        assert "render_results_toolbar" in source
        assert "Export is unavailable" in source
